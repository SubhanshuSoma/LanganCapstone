"""Local document extraction and paragraph-aware chunking for ingestion."""

from __future__ import annotations

import csv
import html
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from email import policy
from email.parser import BytesParser
from pathlib import Path

import tiktoken

ENCODING = tiktoken.get_encoding("cl100k_base")
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".csv", ".eml"}


@dataclass
class Block:
    kind: str
    text: str
    level: int | None = None
    page: int | None = None


@dataclass
class ExtractedDoc:
    blocks: list[Block]
    metadata: dict


@dataclass
class Section:
    heading: str
    blocks: list[Block]
    page_start: int | None
    page_end: int | None

    @property
    def content(self) -> str:
        return "\n\n".join(block.text for block in self.blocks)


@dataclass
class Chunk:
    content: str
    page_start: int | None
    page_end: int | None


def token_count(text: str) -> int:
    return len(ENCODING.encode(text))


def _table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(map(len, rows))
    normalized = [[str(cell or "").replace("|", "\\|").replace("\n", " ").strip() for cell in row] for row in rows]
    normalized = [row + [""] * (width - len(row)) for row in normalized]
    return "\n".join(
        ["| " + " | ".join(normalized[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
        + ["| " + " | ".join(row) + " |" for row in normalized[1:]]
    )


def extract(path: Path) -> ExtractedDoc:
    ext = path.suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")
    if ext in {".txt", ".md"}:
        blocks = []
        for paragraph in re.split(r"\n\s*\n", path.read_text(encoding="utf-8-sig")):
            paragraph = paragraph.strip()
            if not paragraph:
                continue
            for line in paragraph.splitlines():
                match = re.match(r"^(#{1,6})\s+(.+)$", line) if ext == ".md" else None
                if match:
                    blocks.append(Block("heading", match.group(2), len(match.group(1))))
                else:
                    blocks.append(Block("text", line))
        return ExtractedDoc(blocks, {})
    if ext == ".pdf":
        import pdfplumber

        blocks = []
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                for paragraph in re.split(r"\n\s*\n", page.extract_text() or ""):
                    if paragraph.strip():
                        blocks.append(Block("text", paragraph.strip(), page=page_number))
        if sum(len(block.text) for block in blocks) < 50:
            raise ValueError("No text layer found; scanned PDF needs OCR (not supported in Phase 1)")
        return ExtractedDoc(blocks, {})
    if ext == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(path)
        blocks = []
        for item in doc.iter_inner_content():
            if isinstance(item, Paragraph):
                value = item.text.strip()
                if value:
                    match = re.fullmatch(r"Heading ([1-6])", item.style.name or "")
                    blocks.append(Block("heading" if match else "text", value, int(match.group(1)) if match else None))
            elif isinstance(item, Table):
                rows = [[cell.text for cell in row.cells] for row in item.rows]
                if rows:
                    blocks.append(Block("table", _table(rows)))
        return ExtractedDoc(blocks, {})
    if ext == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = [row for _, row in zip(range(2001), csv.reader(stream))]
            truncated = bool(stream.readline())
        return ExtractedDoc([Block("table", _table(rows))] if rows else [], {"truncated": truncated})
    if ext == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks = []
        truncated_sheets = []
        try:
            for sheet in workbook:
                rows = [["" if cell is None else str(cell) for cell in row] for _, row in zip(range(2001), sheet.values)]
                rows = [row for row in rows if any(cell.strip() for cell in row)]
                if rows:
                    blocks.extend([Block("heading", sheet.title, 1), Block("table", _table(rows))])
                if sheet.max_row and sheet.max_row > 2000:
                    truncated_sheets.append(sheet.title)
        finally:
            workbook.close()
        return ExtractedDoc(blocks, {"truncated_sheets": truncated_sheets})
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    metadata = {key.lower(): str(message.get(key, "")) for key in ("Subject", "From", "To", "Date")}
    metadata["attachments"] = [part.get_filename() for part in message.walk() if part.get_filename()]
    plain = message.get_body(preferencelist=("plain",))
    body = plain.get_content() if plain else ""
    if not body:
        html_part = message.get_body(preferencelist=("html",))
        body = re.sub(r"<[^>]+>", " ", html_part.get_content()) if html_part else ""
        body = html.unescape(body)
    body = re.split(r"(?m)^On .+ wrote:\s*$", body)[0]
    body = "\n".join(line for line in body.splitlines() if not line.lstrip().startswith(">"))
    return ExtractedDoc([Block("text", part.strip()) for part in re.split(r"\n\s*\n", body) if part.strip()], metadata)


def clean(document: ExtractedDoc, *, is_pdf: bool = False) -> ExtractedDoc:
    repeated = Counter()
    if is_pdf:
        pages = {block.page for block in document.blocks if block.page is not None}
        for page in pages:
            repeated.update({line.strip() for block in document.blocks if block.page == page for line in block.text.splitlines() if line.strip()})
        repeated = Counter({line: count for line, count in repeated.items() if count > len(pages) / 2 and len(pages) > 1})
    blocks = []
    for block in document.blocks:
        lines = [line for line in block.text.splitlines() if line.strip() not in repeated]
        value = "\n".join(re.sub(r"[ \t]+", " ", unicodedata.normalize("NFKC", line)).strip() for line in lines).strip()
        if value:
            blocks.append(Block(block.kind, value, block.level, block.page))
    return ExtractedDoc(blocks, document.metadata)


def _pages(blocks: list[Block]) -> tuple[int | None, int | None]:
    pages = [block.page for block in blocks if block.page is not None]
    return (min(pages), max(pages)) if pages else (None, None)


def make_sections(blocks: list[Block], max_tokens: int = 1500) -> list[Section]:
    if max_tokens < 1:
        raise ValueError("max_tokens must be positive")
    sections: list[Section] = []
    headings: list[str] = []
    pending: list[Block] = []

    def flush() -> None:
        nonlocal pending
        if pending:
            start, end = _pages(pending)
            sections.append(Section(" > ".join(headings) or "Document", pending, start, end))
            pending = []

    for block in blocks:
        if block.kind == "heading":
            flush()
            level = block.level or 1
            headings[:] = headings[: level - 1] + [block.text]
            continue
        if block.kind == "table" and token_count(block.text) > max_tokens:
            lines = block.text.splitlines()
            header = "\n".join(lines[:2])
            parts = [header + "\n" + line for line in lines[2:]] or [header]
        else:
            parts = _split_text(block.text, max_tokens)
        for part in parts:
            item = Block(block.kind, part, block.level, block.page)
            if pending and token_count("\n\n".join(old.text for old in pending + [item])) > max_tokens:
                flush()
            pending.append(item)
    flush()
    return sections


def _split_text(text: str, max_tokens: int) -> list[str]:
    if token_count(text) <= max_tokens:
        return [text]
    pieces = re.split(r"(?<=[.!?])\s+", text)
    result = []
    current = ""
    for piece in pieces:
        if token_count(piece) > max_tokens:
            words = piece.split()
            for word in words:
                candidate = f"{current} {word}".strip()
                if current and token_count(candidate) > max_tokens:
                    result.append(current)
                    current = word
                else:
                    current = candidate
        else:
            candidate = f"{current} {piece}".strip()
            if current and token_count(candidate) > max_tokens:
                result.append(current)
                current = piece
            else:
                current = candidate
    if current:
        result.append(current)
    return result


def make_chunks(section: Section, max_tokens: int = 300, overlap_tokens: int = 50) -> list[Chunk]:
    if max_tokens < 1 or not 0 <= overlap_tokens < max_tokens:
        raise ValueError("Require max_tokens > overlap_tokens >= 0")
    units: list[Block] = []
    for block in section.blocks:
        if block.kind == "table":
            lines = block.text.splitlines()
            if len(lines) >= 2:
                header = "\n".join(lines[:2])
                units.extend(Block("table", header + "\n" + line, page=block.page) for line in lines[2:])
                if len(lines) == 2:
                    units.append(Block("table", header, page=block.page))
            else:
                units.append(block)
        else:
            units.extend(Block("text", part, page=block.page) for part in _split_text(block.text, max_tokens))
    chunks = []
    current: list[Block] = []
    for unit in units:
        candidate = "\n\n".join(item.text for item in current + [unit])
        if current and token_count(candidate) > max_tokens:
            start, end = _pages(current)
            chunks.append(Chunk("\n\n".join(item.text for item in current), start, end))
            overlap: list[Block] = []
            if overlap_tokens and unit.kind != "table":
                for old in reversed(current):
                    if old.kind == "table" or token_count("\n\n".join(item.text for item in [old] + overlap + [unit])) > max_tokens:
                        break
                    overlap.insert(0, old)
                    if token_count("\n\n".join(item.text for item in overlap)) >= overlap_tokens:
                        break
            current = overlap
        current.append(unit)
    if current:
        start, end = _pages(current)
        chunks.append(Chunk("\n\n".join(item.text for item in current), start, end))
    return chunks


def serialize_blocks(blocks: list[Block]) -> list[dict]:
    return [asdict(block) for block in blocks]
