from pathlib import Path

from docx import Document

from app.services.document_processing import (
    Block,
    clean,
    extract,
    make_chunks,
    make_sections,
    token_count,
)
from scripts.ingest_documents import fake_embedding, sha256_file


def test_docx_headings_sections_and_chunks(tmp_path: Path):
    source = tmp_path / "notes.docx"
    document = Document()
    document.add_heading("Foundations", 1)
    document.add_paragraph("The foundations were installed in phases. " * 15)
    document.add_heading("Piles", 2)
    document.add_paragraph("Pile installation used driven steel sections. " * 15)
    document.save(source)

    extracted = clean(extract(source))
    sections = make_sections(extracted.blocks, max_tokens=1500)
    assert [section.heading for section in sections] == ["Foundations", "Foundations > Piles"]
    chunks = [chunk for section in sections for chunk in make_chunks(section, max_tokens=40, overlap_tokens=8)]
    assert chunks
    assert all(token_count(chunk.content) <= 40 for chunk in chunks)
    assert fake_embedding(chunks[0].content, 16) == fake_embedding(chunks[0].content, 16)
    assert len(fake_embedding(chunks[0].content, 16)) == 16
    assert len(sha256_file(source)) == 64


def test_csv_table_rows_stay_intact(tmp_path: Path):
    source = tmp_path / "data.csv"
    source.write_text("name,value\nA,1\nB,2\nC,3\n", encoding="utf-8")
    sections = make_sections(extract(source).blocks)
    chunks = make_chunks(sections[0], max_tokens=20, overlap_tokens=0)
    assert all("| name | value |" in chunk.content for chunk in chunks)
    assert sum(chunk.content.count("| A | 1 |" ) for chunk in chunks) == 1
    assert sum(chunk.content.count("| B | 2 |" ) for chunk in chunks) == 1


def test_long_paragraph_is_split_into_bounded_sections():
    blocks = [Block("text", "A short synthetic sentence. " * 50)]
    sections = make_sections(blocks, max_tokens=25)
    assert len(sections) > 1
    assert all(token_count(section.content) <= 25 for section in sections)


def test_pdf_repeated_header_removed():
    from app.services.document_processing import ExtractedDoc

    doc = ExtractedDoc([
        Block("text", "Repeated header\nUnique page one", page=1),
        Block("text", "Repeated header\nUnique page two", page=2),
        Block("text", "Repeated header\nUnique page three", page=3),
    ], {})
    cleaned = clean(doc, is_pdf=True)
    assert all("Repeated header" not in block.text for block in cleaned.blocks)
    assert [block.page for block in cleaned.blocks] == [1, 2, 3]
